from __future__ import annotations

from collections.abc import Iterable
import csv
from datetime import date, datetime
from decimal import Decimal
from decimal import InvalidOperation
from io import BytesIO
from io import StringIO
from pathlib import Path
import re
from time import monotonic
from typing import Any
import unicodedata
from uuid import NAMESPACE_URL
from uuid import UUID
from uuid import uuid5
from uuid import uuid4

from fastapi import UploadFile
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.domain import Claim
from app.models.domain import ClaimDocument
from app.models.domain import Insured
from app.models.domain import Policy
from app.models.domain import Provider
from app.models.domain import Vehicle
from app.schemas.claims import ClaimCreate
from app.schemas.claims import ClaimDocumentCreate
from app.schemas.claims import InsuredBase
from app.schemas.claims import PolicyBase
from app.schemas.claims import ProviderBase
from app.schemas.claims import VehicleBase
from app.schemas.imports import DataImportPayload
from app.schemas.imports import FileImportResponse
from app.services.import_service import ImportService

SUPPORTED_DATASETS = {
    "asegurados": "insureds",
    "insureds": "insureds",
    "polizas": "policies",
    "policies": "policies",
    "proveedores": "providers",
    "providers": "providers",
    "vehiculos": "vehicles",
    "vehicles": "vehicles",
    "siniestros": "claims",
    "claims": "claims",
    "documentos": "documents",
    "documents": "documents",
}

DATASET_ORDER = ("insureds", "providers", "policies", "vehicles", "claims", "documents")

DATASET_LABELS = {
    "insureds": "asegurados",
    "providers": "proveedores",
    "policies": "polizas",
    "vehicles": "vehiculos",
    "claims": "siniestros",
    "documents": "documentos",
}

DATASET_REQUIRED_COLUMNS = {
    "insureds": {"id_asegurado"},
    "providers": {"id_proveedor"},
    "policies": {"id_poliza", "id_asegurado", "ramo", "fecha_inicio", "fecha_fin"},
    "vehicles": {"id_vehiculo", "id_poliza"},
    "claims": {"id_siniestro", "id_poliza", "id_asegurado"},
    "documents": {"id_documento", "id_siniestro"},
}

DATASET_COLUMN_ALIASES = {
    "insureds": {
        "id_asegurado": {"id_asegurado", "id"},
        "code": {"code", "codigo", "codigo_asegurado"},
        "nombres_asegurado": {"nombres_asegurado", "nombre_asegurado", "nombres", "nombre"},
        "segmento": {"segmento", "segment"},
        "antiguedad_meses": {"antiguedad_meses", "seniority_months"},
        "antiguedad_anos": {"antiguedad_anos", "antiguedad_anios", "seniority_years"},
        "ciudad": {"ciudad", "city"},
        "num_polizas": {
            "num_polizas",
            "policy_count",
            "n_polizas_activas",
            "numero_polizas_activas",
            "polizas_activas",
        },
        "reclamos_12m": {
            "reclamos_12m",
            "claims_12m",
            "n_reclamos_ultimos_12_meses",
            "numero_reclamos_ultimos_12_meses",
            "reclamos_ultimos_12_meses",
        },
        "mora_actual": {"mora_actual", "current_delinquency"},
        "score_cliente": {"score_cliente", "client_score"},
        "reclamos_historico_total": {
            "reclamos_historico_total",
            "n_reclamos_historico_total",
            "numero_reclamos_historico_total",
        },
        "reclamos_rc_sin_tercero": {"reclamos_rc_sin_tercero"},
        "perfil_riesgo_historico": {"perfil_riesgo_historico"},
    },
    "providers": {
        "id_proveedor": {"id_proveedor", "id"},
        "code": {"code", "codigo", "codigo_proveedor"},
        "nombre": {"nombre", "nombre_proveedor", "name"},
        "tipo": {"tipo", "provider_type"},
        "ciudad": {"ciudad", "city"},
        "reclamos_asociados": {
            "reclamos_asociados",
            "associated_claims",
            "n_siniestros_asociados",
            "numero_siniestros_asociados",
            "siniestros_asociados",
        },
        "monto_promedio": {"monto_promedio", "promedio_monto", "average_amount"},
        "pct_casos_observados": {"pct_casos_observados", "observed_cases_pct"},
        "antiguedad_meses": {"antiguedad_meses", "seniority_months"},
        "en_lista_restrictiva": {"en_lista_restrictiva", "is_restricted"},
        "motivo_restriccion": {"motivo_restriccion", "restriction_reason"},
    },
    "policies": {
        "id_poliza": {"id_poliza", "id"},
        "code": {"code", "codigo", "codigo_poliza"},
        "id_asegurado": {"id_asegurado", "insured_id"},
        "ramo": {"ramo", "branch"},
        "fecha_inicio": {"fecha_inicio", "start_date"},
        "fecha_fin": {"fecha_fin", "end_date"},
        "prima": {"prima", "prima_anual", "premium_amount"},
        "suma_asegurada": {"suma_asegurada", "insured_amount"},
        "deducible": {"deducible", "deductible"},
        "canal_venta": {"canal_venta", "sales_channel"},
        "ciudad": {"ciudad", "city"},
        "estado_poliza": {"estado_poliza", "status"},
    },
    "vehicles": {
        "id_vehiculo": {"id_vehiculo", "id"},
        "code": {"code", "codigo", "codigo_vehiculo"},
        "id_poliza": {"id_poliza", "policy_id"},
        "id_asegurado": {"id_asegurado", "insured_id"},
        "placa": {"placa", "plate"},
        "chasis": {"chasis", "chassis"},
        "motor": {"motor", "engine"},
        "marca": {"marca", "brand"},
        "modelo": {"modelo", "model"},
        "anio": {"anio", "year"},
        "color": {"color"},
    },
    "claims": {
        "id_siniestro": {"id_siniestro", "id"},
        "code": {"code", "codigo", "codigo_siniestro"},
        "id_poliza": {"id_poliza", "policy_id"},
        "id_asegurado": {"id_asegurado", "insured_id"},
        "id_proveedor": {"id_proveedor", "provider_id"},
        "id_vehiculo": {"id_vehiculo", "vehicle_id"},
        "placa_vehiculo_asegurado": {"placa_vehiculo_asegurado", "placa", "vehicle_plate"},
        "ramo": {"ramo", "branch"},
        "cobertura": {"cobertura", "coverage"},
        "fecha_ocurrencia": {"fecha_ocurrencia", "occurrence_date"},
        "fecha_reporte": {"fecha_reporte", "reported_date"},
        "monto_reclamado": {"monto_reclamado", "claimed_amount"},
        "monto_estimado": {"monto_estimado", "estimated_amount"},
        "monto_pagado": {"monto_pagado", "paid_amount"},
        "estado": {"estado", "status"},
        "sucursal": {"sucursal", "office"},
        "descripcion": {"descripcion", "descripcion_del_evento", "description"},
        "documentos_completos": {"documentos_completos", "docs_completos", "documents_complete"},
        "dias_desde_inicio_poliza": {"dias_desde_inicio_poliza", "days_from_policy_start"},
        "dias_desde_fin_poliza": {
            "dias_desde_fin_poliza",
            "dias_hasta_fin_poliza",
            "days_from_policy_end",
        },
        "dias_entre_ocurrencia_reporte": {
            "dias_entre_ocurrencia_reporte",
            "dias_ocurr_reporte",
            "report_delay_days",
        },
        "historial_siniestros_asegurado": {
            "historial_siniestros_asegurado",
            "n_reclamos_previos_asegurado",
            "numero_reclamos_previos_asegurado",
            "insured_claim_history",
        },
        "proveedor_lista_restrictiva": {
            "proveedor_lista_restrictiva",
            "prov_lista_restrictiva",
            "provider_restricted",
        },
        "suma_asegurada": {"suma_asegurada", "insured_amount"},
        "ratio_monto_suma_asegurada": {"ratio_monto_suma_asegurada", "amount_to_insured_ratio"},
        "similitud_narrativa_max": {"similitud_narrativa_max", "narrative_similarity_max"},
        "numero_parte_policial": {"numero_parte_policial", "police_report_number"},
        "estado_flujo": {"estado_flujo", "workflow_status"},
        "ultima_decision": {"ultima_decision", "last_decision"},
        "ultima_revision_en": {"ultima_revision_en", "last_review_at"},
    },
    "documents": {
        "id_documento": {"id_documento", "id"},
        "id_siniestro": {"id_siniestro", "claim_id"},
        "tipo_documento": {"tipo_documento", "document_type"},
        "entregado": {"entregado", "delivered"},
        "legible": {"legible"},
        "fecha_emision": {"fecha_emision", "issue_date"},
        "inconsistencia_detectada": {"inconsistencia_detectada", "inconsistency_detected"},
        "observacion": {"observacion", "notes"},
        "nombre_archivo_pdf": {"nombre_archivo_pdf", "file_name_pdf"},
    },
}

DATASET_ALLOWED_COLUMNS = {
    dataset: {alias for aliases in columns.values() for alias in aliases}
    for dataset, columns in DATASET_COLUMN_ALIASES.items()
}

DATASET_SIGNATURE_COLUMN_KEYS = {
    "insureds": {
        "nombres_asegurado",
        "segmento",
        "antiguedad_anos",
        "num_polizas",
        "reclamos_12m",
        "mora_actual",
        "score_cliente",
        "reclamos_historico_total",
        "reclamos_rc_sin_tercero",
        "perfil_riesgo_historico",
    },
    "providers": {
        "nombre",
        "tipo",
        "reclamos_asociados",
        "monto_promedio",
        "pct_casos_observados",
        "en_lista_restrictiva",
        "motivo_restriccion",
    },
    "policies": {
        "fecha_inicio",
        "fecha_fin",
        "prima",
        "suma_asegurada",
        "deducible",
        "canal_venta",
        "estado_poliza",
    },
    "vehicles": {"id_vehiculo", "id_asegurado", "placa", "chasis", "motor", "marca", "modelo", "anio", "color"},
    "claims": {
        "id_siniestro",
        "id_vehiculo",
        "placa_vehiculo_asegurado",
        "cobertura",
        "fecha_ocurrencia",
        "fecha_reporte",
        "monto_reclamado",
        "monto_estimado",
        "monto_pagado",
        "estado",
        "sucursal",
        "descripcion",
        "documentos_completos",
        "dias_desde_inicio_poliza",
        "dias_desde_fin_poliza",
        "dias_entre_ocurrencia_reporte",
        "historial_siniestros_asegurado",
        "proveedor_lista_restrictiva",
        "suma_asegurada",
        "ratio_monto_suma_asegurada",
        "similitud_narrativa_max",
        "numero_parte_policial",
        "estado_flujo",
        "ultima_decision",
        "ultima_revision_en",
    },
    "documents": {
        "id_documento",
        "tipo_documento",
        "entregado",
        "legible",
        "fecha_emision",
        "inconsistencia_detectada",
        "observacion",
        "nombre_archivo_pdf",
    },
}

DATASET_SIGNATURE_COLUMNS = {
    dataset: {
        alias
        for canonical in columns
        for alias in DATASET_COLUMN_ALIASES[dataset].get(canonical, {canonical})
    }
    for dataset, columns in DATASET_SIGNATURE_COLUMN_KEYS.items()
}

DATASET_REQUIRED_COLUMN_GROUPS = {
    dataset: {
        canonical: DATASET_COLUMN_ALIASES[dataset].get(canonical, {canonical})
        for canonical in required_columns
    }
    for dataset, required_columns in DATASET_REQUIRED_COLUMNS.items()
}

BRANCH_NORMALIZATION = {
    "vida": "Vida",
    "salud": "Salud",
    "vehiculo": "Vehiculos",
    "vehiculos": "Vehiculos",
    "hogar": "Hogar",
    "general": "Generales",
    "generales": "Generales",
}

CLAIM_STATUS_NORMALIZATION = {
    "abierto": "Abierto",
    "analisis": "En analisis",
    "en_analisis": "En analisis",
    "revision": "En revision",
    "en_revision": "En revision",
    "observado": "Observado",
    "investigacion": "Investigacion",
    "pendiente": "Pendiente",
    "reserva": "Reserva",
    "cerrado": "Cerrado",
    "finalizado": "Finalizado",
    "pagado": "Pagado",
    "pago_total": "Pago Total",
    "pago_parcial": "Pago Parcial",
    "rechazado": "Rechazado",
    "anulado": "Anulado",
    "cancelado": "Cancelado",
}

POLICY_STATUS_NORMALIZATION = {
    "activa": "Vigente",
    "activo": "Vigente",
    "vigente": "Vigente",
    "cancelada": "Cancelada",
    "cancelado": "Cancelada",
    "vencida": "Vencida",
    "vencido": "Vencida",
    "expirada": "Expirada",
    "expirado": "Expirada",
}


class _ImportTimeoutGuard:
    def __init__(self, timeout_seconds: float) -> None:
        self.timeout_seconds = timeout_seconds
        self.deadline = monotonic() + timeout_seconds if timeout_seconds > 0 else None

    def __call__(self) -> None:
        if self.deadline is not None and monotonic() > self.deadline:
            raise TimeoutError(
                f"La importacion supero el timeout configurado de {self.timeout_seconds:.0f} segundos."
            )


class FileImportService:
    def __init__(self, import_service: ImportService | None = None) -> None:
        self.import_service = import_service or ImportService()

    def import_file(
        self,
        db: Session,
        file: UploadFile,
        *,
        dataset: str | None = None,
        reset: bool = False,
        recalculate_scores: bool = True,
    ) -> FileImportResponse:
        guard = _ImportTimeoutGuard(settings.import_timeout_seconds)
        guard()

        content = file.file.read()
        filename = file.filename or "archivo"
        suffix = Path(filename).suffix.lower()

        if suffix == ".csv":
            rows_by_dataset = self._read_csv(content, filename=filename, dataset=dataset)
        elif suffix in {".xlsx", ".xlsm"}:
            rows_by_dataset = self._read_excel(content, dataset=dataset)
        else:
            raise ValueError("Formato no soportado. Usa .csv, .xlsx o .xlsm.")

        self._validate_row_limit(rows_by_dataset)
        guard()

        payload, standalone_documents, warnings, skipped_rows = self._payload_from_rows(db, rows_by_dataset)
        result = self.import_service.import_payload(
            db,
            payload,
            reset=reset,
            assess_claims=recalculate_scores,
            use_embeddings=False,
            should_continue=guard,
        )
        document_claim_ids = self._import_standalone_documents(db, standalone_documents)
        if recalculate_scores and document_claim_ids:
            result["assessments"] += self.import_service.risk_service.assess_claims(
                db,
                sorted(document_claim_ids),
                use_embeddings=False,
                should_continue=guard,
            )

        return FileImportResponse(
            message="Archivo importado correctamente",
            filename=filename,
            datasets={name: len(rows_by_dataset.get(name, [])) for name in DATASET_ORDER},
            warnings=warnings,
            skipped_rows=skipped_rows,
            **result,
        )

    def _read_csv(self, content: bytes, *, filename: str, dataset: str | None) -> dict[str, list[dict[str, Any]]]:
        filename_dataset = self._resolve_dataset(Path(filename).stem, required=False)

        text = self._decode_text(content)
        sample = text[:2048]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;|\t") if sample.strip() else csv.excel
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(StringIO(text), dialect=dialect)
        headers = {self._normalize_header(header) for header in reader.fieldnames or [] if header}
        dataset_key = self._resolve_csv_dataset(
            dataset=dataset,
            filename=filename,
            filename_dataset=filename_dataset,
            headers=headers,
        )
        self._validate_headers(dataset_key, headers, source=filename)
        rows = [cleaned for row in reader if (cleaned := self._clean_row(row))]
        return {dataset_key: rows}

    def _resolve_csv_dataset(
        self,
        *,
        dataset: str | None,
        filename: str,
        filename_dataset: str | None,
        headers: set[str],
    ) -> str:
        if dataset:
            dataset_key = self._resolve_dataset(dataset)
            if filename_dataset and filename_dataset != dataset_key:
                raise ValueError(
                    f"El archivo '{filename}' corresponde a {DATASET_LABELS[filename_dataset]}, "
                    f"pero seleccionaste dataset={DATASET_LABELS[dataset_key]}. "
                    "Corrige el dataset o cambia el nombre del archivo."
                )
            return dataset_key

        if filename_dataset:
            return filename_dataset

        return self._infer_dataset_from_headers(headers, source=filename)

    def _infer_dataset_from_headers(self, headers: set[str], *, source: str) -> str:
        headers = {header for header in headers if header}
        if not headers:
            raise ValueError(f"{source}: el archivo no tiene encabezados.")

        candidates: list[tuple[int, str]] = []
        for dataset in DATASET_ORDER:
            missing = self._missing_required_headers(dataset, headers)
            incompatible = self._incompatible_header_messages(dataset, headers)
            if missing or incompatible:
                continue
            score = len(headers & DATASET_ALLOWED_COLUMNS[dataset]) + len(headers & DATASET_SIGNATURE_COLUMNS[dataset])
            candidates.append((score, dataset))

        if not candidates:
            allowed = ", ".join(DATASET_LABELS[dataset] for dataset in DATASET_ORDER)
            raise ValueError(
                f"{source}: no se pudo detectar automaticamente el dataset por nombre ni por columnas. "
                f"Usa un nombre de archivo como {allowed} o envia encabezados reconocibles. "
                f"Columnas recibidas: {', '.join(sorted(headers))}."
            )

        candidates.sort(reverse=True)
        best_score, dataset = candidates[0]
        tied = [candidate for score, candidate in candidates if score == best_score]
        if len(tied) > 1:
            labels = ", ".join(DATASET_LABELS[candidate] for candidate in tied)
            raise ValueError(
                f"{source}: columnas ambiguas para detectar dataset automaticamente ({labels}). "
                "Usa un nombre de archivo reconocido o encabezados mas especificos."
            )

        return dataset

    def _validate_row_limit(self, rows_by_dataset: dict[str, list[dict[str, Any]]]) -> None:
        total_rows = sum(len(rows) for rows in rows_by_dataset.values())
        if settings.import_max_rows > 0 and total_rows > settings.import_max_rows:
            raise ValueError(
                f"El archivo contiene {total_rows} filas importables y supera el limite "
                f"configurado de {settings.import_max_rows}. Divide la carga en lotes mas pequenos."
            )

    def _read_excel(self, content: bytes, *, dataset: str | None) -> dict[str, list[dict[str, Any]]]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Para importar Excel instala la dependencia openpyxl.") from exc

        workbook = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
        rows_by_dataset: dict[str, list[dict[str, Any]]] = {}

        for sheet in workbook.worksheets:
            sheet_dataset = self._resolve_dataset(sheet.title, required=False)
            dataset_key = sheet_dataset
            if dataset_key is None and dataset:
                dataset_key = self._resolve_dataset(dataset)
            elif dataset and sheet_dataset and sheet_dataset != self._resolve_dataset(dataset):
                raise ValueError(
                    f"La hoja '{sheet.title}' corresponde a {DATASET_LABELS[sheet_dataset]}, "
                    f"pero seleccionaste dataset={DATASET_LABELS[self._resolve_dataset(dataset)]}."
                )
            if dataset_key is None:
                continue

            rows = list(self._sheet_rows(sheet.iter_rows(values_only=True), dataset=dataset_key, source=sheet.title))
            if rows:
                rows_by_dataset.setdefault(dataset_key, []).extend(rows)

        if not rows_by_dataset:
            raise ValueError(
                "No encontre hojas importables. Usa nombres como asegurados, polizas, proveedores, "
                "vehiculos, siniestros o documentos."
            )
        return rows_by_dataset

    def _payload_from_rows(
        self,
        db: Session,
        rows_by_dataset: dict[str, list[dict[str, Any]]],
    ) -> tuple[DataImportPayload, list[tuple[str, ClaimDocumentCreate]], list[str], int]:
        payload = DataImportPayload()
        documents_by_claim: dict[str, list[ClaimDocumentCreate]] = {}
        standalone_documents: list[tuple[str, ClaimDocumentCreate]] = []
        identities = self._load_identity_maps(db)
        batch_code_owners: dict[str, dict[str, str]] = {dataset: {} for dataset in identities}

        errors: list[str] = []
        warnings: list[str] = []
        skipped_rows = 0
        skipped_document_claim_codes: list[str] = []
        for dataset in DATASET_ORDER:
            rows = rows_by_dataset.get(dataset, [])
            for index, row in enumerate(rows, start=2):
                try:
                    if dataset == "insureds":
                        payload.insureds.append(InsuredBase(**self._map_insured(row, identities, batch_code_owners)))
                    elif dataset == "providers":
                        payload.providers.append(ProviderBase(**self._map_provider(row, identities, batch_code_owners)))
                    elif dataset == "policies":
                        payload.policies.append(PolicyBase(**self._map_policy(row, identities, batch_code_owners)))
                    elif dataset == "vehicles":
                        payload.vehicles.append(VehicleBase(**self._map_vehicle(row, identities, batch_code_owners)))
                    elif dataset == "claims":
                        payload.claims.append(ClaimCreate(**self._map_claim(row, identities, batch_code_owners)))
                    elif dataset == "documents":
                        document = ClaimDocumentCreate(**self._map_document(row))
                        raw_claim_id = self._value(row, "id_siniestro", "claim_id")
                        if not raw_claim_id:
                            raise ValueError("documentos requiere id_siniestro")
                        claim_id = self._resolve_reference(
                            row,
                            "claims",
                            identities,
                            "id_siniestro",
                            "claim_id",
                            required=False,
                        )
                        if not claim_id:
                            skipped_rows += 1
                            if len(skipped_document_claim_codes) < 5:
                                skipped_document_claim_codes.append(self._clean_code(raw_claim_id) or str(raw_claim_id))
                            continue
                        documents_by_claim.setdefault(str(claim_id), []).append(document)
                except (ValidationError, ValueError, TypeError) as exc:
                    errors.append(f"{dataset} fila {index}: {exc}")

        if errors:
            preview = "; ".join(errors[:5])
            raise ValueError(f"Archivo invalido. {preview}")

        if documents_by_claim:
            claim_ids_in_payload = {claim.id for claim in payload.claims}
            for claim in payload.claims:
                claim.documents.extend(documents_by_claim.pop(claim.id, []))
            if documents_by_claim:
                for claim_id, documents in documents_by_claim.items():
                    if claim_id not in claim_ids_in_payload:
                        standalone_documents.extend((claim_id, document) for document in documents)

        if skipped_rows:
            examples = ", ".join(dict.fromkeys(skipped_document_claim_codes))
            examples_text = f" (ejemplos: {examples})" if examples else ""
            warnings.append(
                f"Se omitieron {skipped_rows} documentos porque referencian siniestros no cargados"
                f"{examples_text}."
            )

        return payload, standalone_documents, warnings, skipped_rows

    def _import_standalone_documents(
        self,
        db: Session,
        documents: list[tuple[str, ClaimDocumentCreate]],
    ) -> set[str]:
        claim_ids: set[str] = set()
        for claim_id, document in documents:
            db.merge(ClaimDocument(claim_id=claim_id, **document.model_dump()))
            claim_ids.add(claim_id)
        if documents:
            db.commit()
        return claim_ids

    def _load_identity_maps(self, db: Session) -> dict[str, dict[str, str]]:
        models = {
            "insureds": Insured,
            "providers": Provider,
            "policies": Policy,
            "vehicles": Vehicle,
            "claims": Claim,
        }
        identity_maps: dict[str, dict[str, str]] = {dataset: {} for dataset in models}
        for dataset, model in models.items():
            rows = db.execute(
                select(model.id, model.code).where(model.code.is_not(None))
            ).all()
            identity_maps[dataset] = {
                clean_code: str(record_id)
                for record_id, code in rows
                if (clean_code := self._clean_code(code))
            }
        for record_id, plate in db.execute(select(Vehicle.id, Vehicle.plate).where(Vehicle.plate.is_not(None))).all():
            if clean_plate := self._clean_code(plate):
                identity_maps["vehicles"][clean_plate] = str(record_id)
        return identity_maps

    def _record_identity(
        self,
        row: dict[str, Any],
        dataset: str,
        identities: dict[str, dict[str, str]],
        batch_code_owners: dict[str, dict[str, str]],
        *,
        id_keys: tuple[str, ...],
        code_keys: tuple[str, ...],
    ) -> tuple[str, str | None]:
        raw_id = self._value(row, *id_keys)
        explicit_code = self._value(row, *code_keys)
        raw_id_text = self._clean_text(raw_id)
        code = self._clean_code(explicit_code)
        raw_id_is_uuid = bool(raw_id_text and self._is_uuid(raw_id_text))

        if raw_id_is_uuid:
            record_id = raw_id_text
        else:
            code = code or self._clean_code(raw_id_text)
            record_id = identities[dataset].get(code or "")
            if record_id is None:
                record_id = self._stable_id(dataset, code) if code else str(uuid4())

        if code:
            batch_owner = batch_code_owners[dataset].get(code)
            if batch_owner and batch_owner != record_id:
                raise ValueError(f"Codigo duplicado en la carga de {DATASET_LABELS[dataset]}: {code}.")
            existing_id = identities[dataset].get(code)
            if existing_id and existing_id != record_id:
                if raw_id_is_uuid:
                    raise ValueError(
                        f"Codigo duplicado en {DATASET_LABELS[dataset]}: {code} ya pertenece a {existing_id}."
                    )
                record_id = existing_id
            identities[dataset][code] = record_id
            batch_code_owners[dataset][code] = record_id

        return record_id, code

    def _resolve_reference(
        self,
        row: dict[str, Any],
        dataset: str,
        identities: dict[str, dict[str, str]],
        *keys: str,
        required: bool = True,
    ) -> str | None:
        value = self._value(row, *keys)
        if value is None:
            if required:
                raise ValueError(f"Campo requerido faltante: {keys[0]}")
            return None

        text = self._clean_text(value)
        if not text:
            if required:
                raise ValueError(f"Campo requerido faltante: {keys[0]}")
            return None
        if self._is_uuid(text):
            return text

        code = self._clean_code(text)
        record_id = identities[dataset].get(code or "")
        if record_id:
            return record_id
        if required:
            raise ValueError(
                f"No se encontro {DATASET_LABELS[dataset]} con code={code}. "
                "Importa primero el dataset relacionado o incluye la hoja en el mismo Excel."
            )
        return None

    def _vehicle_id_value(self, row: dict[str, Any], identities: dict[str, dict[str, str]]) -> str | None:
        vehicle_id = self._resolve_reference(
            row,
            "vehicles",
            identities,
            "id_vehiculo",
            "vehicle_id",
            required=False,
        )
        if vehicle_id:
            return vehicle_id

        plate = self._clean_code(self._value(row, "placa_vehiculo_asegurado", "placa", "vehicle_plate"))
        if not plate or plate == "N/A":
            return None
        return identities["vehicles"].get(plate)

    def _document_id_value(self, row: dict[str, Any]) -> str:
        raw_id = self._value(row, "id_documento", "id")
        text = self._clean_text(raw_id)
        if text and self._is_uuid(text):
            return text
        code = self._clean_code(text)
        return self._stable_id("documents", code) if code else str(uuid4())

    def _map_insured(
        self,
        row: dict[str, Any],
        identities: dict[str, dict[str, str]],
        batch_code_owners: dict[str, dict[str, str]],
    ) -> dict[str, Any]:
        record_id, code = self._record_identity(
            row,
            "insureds",
            identities,
            batch_code_owners,
            id_keys=("id_asegurado", "id"),
            code_keys=("code", "codigo", "codigo_asegurado"),
        )
        return {
            "id": record_id,
            "code": code,
            "name": self._value(row, "nombres_asegurado", "nombre_asegurado", "nombres", "nombre"),
            "segment": self._value(row, "segmento", "segment"),
            "seniority_months": self._seniority_months_value(row),
            "city": self._value(row, "ciudad", "city"),
            "policy_count": self._int_value(
                row,
                "num_polizas",
                "policy_count",
                "n_polizas_activas",
                "numero_polizas_activas",
                "polizas_activas",
                default=0,
            ),
            "claims_12m": self._int_value(
                row,
                "reclamos_12m",
                "claims_12m",
                "n_reclamos_ultimos_12_meses",
                "numero_reclamos_ultimos_12_meses",
                "reclamos_ultimos_12_meses",
                default=0,
            ),
            "current_delinquency": self._bool_value(row, "mora_actual", "current_delinquency", default=False),
            "client_score": self._decimal_value(row, "score_cliente", "client_score"),
            "historical_claims_total": self._int_value(
                row,
                "reclamos_historico_total",
                "n_reclamos_historico_total",
                "numero_reclamos_historico_total",
            ),
            "rc_claims_without_third_party": self._int_value(row, "reclamos_rc_sin_tercero"),
            "historical_risk_profile": self._value(row, "perfil_riesgo_historico"),
        }

    def _map_policy(
        self,
        row: dict[str, Any],
        identities: dict[str, dict[str, str]],
        batch_code_owners: dict[str, dict[str, str]],
    ) -> dict[str, Any]:
        record_id, code = self._record_identity(
            row,
            "policies",
            identities,
            batch_code_owners,
            id_keys=("id_poliza", "id"),
            code_keys=("code", "codigo", "codigo_poliza"),
        )
        return {
            "id": record_id,
            "code": code,
            "insured_id": self._resolve_reference(row, "insureds", identities, "id_asegurado", "insured_id"),
            "branch": self._normalized_option(row, BRANCH_NORMALIZATION, "ramo", "branch")
            or self._required(row, "ramo", "branch"),
            "start_date": self._date_value(row, "fecha_inicio", "start_date", required=True),
            "end_date": self._date_value(row, "fecha_fin", "end_date", required=True),
            "premium_amount": self._decimal_value(row, "prima", "prima_anual", "premium_amount"),
            "insured_amount": self._decimal_value(row, "suma_asegurada", "insured_amount"),
            "deductible": self._decimal_value(row, "deducible", "deductible"),
            "sales_channel": self._value(row, "canal_venta", "sales_channel"),
            "city": self._value(row, "ciudad", "city"),
            "status": self._normalized_option(row, POLICY_STATUS_NORMALIZATION, "estado_poliza", "status"),
        }

    def _map_provider(
        self,
        row: dict[str, Any],
        identities: dict[str, dict[str, str]],
        batch_code_owners: dict[str, dict[str, str]],
    ) -> dict[str, Any]:
        record_id, code = self._record_identity(
            row,
            "providers",
            identities,
            batch_code_owners,
            id_keys=("id_proveedor", "id"),
            code_keys=("code", "codigo", "codigo_proveedor"),
        )
        return {
            "id": record_id,
            "code": code,
            "name": self._value(row, "nombre", "nombre_proveedor", "name"),
            "provider_type": self._value(row, "tipo", "provider_type"),
            "city": self._value(row, "ciudad", "city"),
            "associated_claims": self._int_value(
                row,
                "reclamos_asociados",
                "associated_claims",
                "n_siniestros_asociados",
                "numero_siniestros_asociados",
                "siniestros_asociados",
                default=0,
            ),
            "average_amount": self._decimal_value(row, "monto_promedio", "promedio_monto", "average_amount"),
            "observed_cases_pct": self._decimal_value(row, "pct_casos_observados", "observed_cases_pct"),
            "seniority_months": self._int_value(row, "antiguedad_meses", "seniority_months"),
            "is_restricted": self._bool_value(row, "en_lista_restrictiva", "is_restricted", default=False),
            "restriction_reason": self._value(row, "motivo_restriccion", "restriction_reason"),
        }

    def _map_vehicle(
        self,
        row: dict[str, Any],
        identities: dict[str, dict[str, str]],
        batch_code_owners: dict[str, dict[str, str]],
    ) -> dict[str, Any]:
        record_id, code = self._record_identity(
            row,
            "vehicles",
            identities,
            batch_code_owners,
            id_keys=("id_vehiculo", "id"),
            code_keys=("code", "codigo", "codigo_vehiculo"),
        )
        mapped = {
            "id": record_id,
            "code": code,
            "policy_id": self._resolve_reference(row, "policies", identities, "id_poliza", "policy_id"),
            "insured_id": self._resolve_reference(
                row,
                "insureds",
                identities,
                "id_asegurado",
                "insured_id",
                required=False,
            ),
            "plate": self._value(row, "placa", "plate"),
            "chassis": self._value(row, "chasis", "chassis"),
            "engine": self._value(row, "motor", "engine"),
            "brand": self._value(row, "marca", "brand"),
            "model": self._value(row, "modelo", "model"),
            "year": self._int_value(row, "anio", "year"),
            "color": self._value(row, "color"),
        }
        if clean_plate := self._clean_code(mapped["plate"]):
            identities["vehicles"][clean_plate] = record_id
        return mapped

    def _map_claim(
        self,
        row: dict[str, Any],
        identities: dict[str, dict[str, str]],
        batch_code_owners: dict[str, dict[str, str]],
    ) -> dict[str, Any]:
        record_id, code = self._record_identity(
            row,
            "claims",
            identities,
            batch_code_owners,
            id_keys=("id_siniestro", "id"),
            code_keys=("code", "codigo", "codigo_siniestro"),
        )
        policy_insured_amount = self._decimal_value(row, "suma_asegurada", "insured_amount")
        amount_ratio = self._decimal_value(row, "ratio_monto_suma_asegurada", "amount_to_insured_ratio")
        claimed_amount = self._decimal_value(row, "monto_reclamado", "claimed_amount")
        if amount_ratio is None and policy_insured_amount and policy_insured_amount > 0 and claimed_amount is not None:
            amount_ratio = claimed_amount / policy_insured_amount

        return {
            "id": record_id,
            "code": code,
            "policy_id": self._resolve_reference(row, "policies", identities, "id_poliza", "policy_id"),
            "insured_id": self._resolve_reference(row, "insureds", identities, "id_asegurado", "insured_id"),
            "provider_id": self._resolve_reference(
                row,
                "providers",
                identities,
                "id_proveedor",
                "provider_id",
                required=False,
            ),
            "vehicle_id": self._vehicle_id_value(row, identities),
            "branch": self._normalized_option(row, BRANCH_NORMALIZATION, "ramo", "branch"),
            "coverage": self._value(row, "cobertura", "coverage"),
            "occurrence_date": self._date_value(row, "fecha_ocurrencia", "occurrence_date"),
            "reported_date": self._date_value(row, "fecha_reporte", "reported_date"),
            "claimed_amount": claimed_amount,
            "estimated_amount": self._decimal_value(row, "monto_estimado", "estimated_amount"),
            "paid_amount": self._decimal_value(row, "monto_pagado", "paid_amount"),
            "status": self._normalized_option(row, CLAIM_STATUS_NORMALIZATION, "estado", "status"),
            "office": self._value(row, "sucursal", "office"),
            "description": self._value(row, "descripcion", "descripcion_del_evento", "description"),
            "documents_complete": self._bool_value(
                row,
                "documentos_completos",
                "docs_completos",
                "documents_complete",
                default=False,
            ),
            "days_from_policy_start": self._int_value(row, "dias_desde_inicio_poliza", "days_from_policy_start"),
            "days_from_policy_end": self._int_value(
                row,
                "dias_desde_fin_poliza",
                "dias_hasta_fin_poliza",
                "days_from_policy_end",
            ),
            "report_delay_days": self._int_value(
                row,
                "dias_entre_ocurrencia_reporte",
                "dias_ocurr_reporte",
                "report_delay_days",
            ),
            "insured_claim_history": self._int_value(
                row,
                "historial_siniestros_asegurado",
                "n_reclamos_previos_asegurado",
                "numero_reclamos_previos_asegurado",
                "insured_claim_history",
                default=0,
            ),
            "workflow_status": self._value(row, "estado_flujo", "workflow_status"),
            "last_decision": self._value(row, "ultima_decision", "last_decision"),
            "last_review_at": self._datetime_value(row, "ultima_revision_en", "last_review_at"),
            "provider_restricted": self._bool_value(
                row,
                "proveedor_lista_restrictiva",
                "prov_lista_restrictiva",
                "provider_restricted",
                default=False,
            ),
            "narrative_similarity_max": self._decimal_value(
                row,
                "similitud_narrativa_max",
                "narrative_similarity_max",
            ),
            "police_report_number": self._value(row, "numero_parte_policial", "police_report_number"),
            "policy_insured_amount": policy_insured_amount,
            "amount_to_insured_ratio": amount_ratio,
        }

    def _map_document(self, row: dict[str, Any]) -> dict[str, Any]:
        return {
            "id": self._document_id_value(row),
            "document_type": self._value(row, "tipo_documento", "document_type"),
            "delivered": self._bool_value(row, "entregado", "delivered", default=True),
            "legible": self._bool_value(row, "legible", default=True),
            "issue_date": self._date_value(row, "fecha_emision", "issue_date"),
            "inconsistency_detected": self._bool_value(
                row,
                "inconsistencia_detectada",
                "inconsistency_detected",
                default=False,
            ),
            "notes": self._value(row, "observacion", "notes"),
            "file_name_pdf": self._value(row, "nombre_archivo_pdf", "file_name_pdf"),
        }

    def _sheet_rows(
        self,
        rows: Iterable[tuple[Any, ...]],
        *,
        dataset: str,
        source: str,
    ) -> Iterable[dict[str, Any]]:
        iterator = iter(rows)
        headers = next(iterator, None)
        if not headers:
            return
        normalized_headers = [self._normalize_header(header) for header in headers]
        self._validate_headers(dataset, set(normalized_headers), source=f"hoja {source}")
        for values in iterator:
            row = {
                header: value
                for header, value in zip(normalized_headers, values, strict=False)
                if header
            }
            cleaned = self._clean_row(row)
            if cleaned:
                yield cleaned

    def _resolve_dataset(self, raw: str | None, *, required: bool = True) -> str | None:
        key = self._normalize_header(raw)
        dataset = SUPPORTED_DATASETS.get(key)
        if dataset is None:
            for token in key.split("_"):
                if token in SUPPORTED_DATASETS:
                    dataset = SUPPORTED_DATASETS[token]
                    break
        if dataset is None and required:
            allowed = ", ".join(sorted({key for key in SUPPORTED_DATASETS if key in {"asegurados", "polizas", "proveedores", "vehiculos", "siniestros", "documentos"}}))
            raise ValueError(f"Dataset no soportado: {raw}. Usa uno de: {allowed}.")
        return dataset

    def _validate_headers(self, dataset: str, headers: set[str], *, source: str) -> None:
        headers = {header for header in headers if header}
        if not headers:
            raise ValueError(f"{source}: el archivo no tiene encabezados.")

        missing = self._missing_required_headers(dataset, headers)
        incompatible_messages = self._incompatible_header_messages(dataset, headers)

        if missing or incompatible_messages:
            message = (
                f"{source}: columnas no corresponden al dataset {DATASET_LABELS[dataset]}."
            )
            if missing:
                message += f" Faltan obligatorias: {', '.join(sorted(missing))}."
            if incompatible_messages:
                message += f" Se detectaron { '; '.join(incompatible_messages[:3]) }."
            message += f" Columnas recibidas: {', '.join(sorted(headers))}."
            raise ValueError(message)

    def _missing_required_headers(self, dataset: str, headers: set[str]) -> set[str]:
        required_groups = DATASET_REQUIRED_COLUMN_GROUPS[dataset]
        return {
            canonical
            for canonical, aliases in required_groups.items()
            if not headers & aliases
        }

    def _incompatible_header_messages(self, dataset: str, headers: set[str]) -> list[str]:
        incompatible_messages: list[str] = []
        allowed_columns = DATASET_ALLOWED_COLUMNS[dataset]
        for other_dataset, signature_columns in DATASET_SIGNATURE_COLUMNS.items():
            if other_dataset == dataset:
                continue
            present = sorted((headers - allowed_columns) & signature_columns)
            if present:
                incompatible_messages.append(
                    f"columnas de {DATASET_LABELS[other_dataset]}: {', '.join(present[:6])}"
                )
        return incompatible_messages

    def _decode_text(self, content: bytes) -> str:
        for encoding in ("utf-8-sig", "utf-8", "latin-1"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="replace")

    def _clean_row(self, row: dict[str, Any]) -> dict[str, Any]:
        cleaned: dict[str, Any] = {}
        for raw_key, value in row.items():
            key = self._normalize_header(raw_key)
            if not key:
                continue
            cleaned_value = self._clean_cell_value(value)
            if cleaned_value is not None:
                cleaned[key] = cleaned_value
        return cleaned

    def _normalize_header(self, value: Any) -> str:
        if value is None:
            return ""
        normalized = unicodedata.normalize("NFKD", str(value).strip().lower())
        normalized = "".join(char for char in normalized if not unicodedata.combining(char))
        normalized = re.sub(r"[^a-z0-9]+", "_", normalized)
        return re.sub(r"_+", "_", normalized).strip("_")

    def _clean_cell_value(self, value: Any) -> Any:
        if not isinstance(value, str):
            return value
        text = value.strip()
        if self._is_blank_marker(text):
            return None
        return text

    def _clean_text(self, value: Any) -> str | None:
        if value is None:
            return None
        text = str(value).strip()
        return None if self._is_blank_marker(text) else text

    def _is_blank_marker(self, value: str) -> bool:
        return value.strip().lower() in {"", "-", "--", "—", "n/a", "na", "null", "none", "sin dato"}

    def _value(self, row: dict[str, Any], *keys: str) -> Any:
        for key in keys:
            normalized = self._normalize_header(key)
            if normalized in row and row[normalized] is not None:
                return row[normalized]
        return None

    def _normalized_option(self, row: dict[str, Any], mapping: dict[str, str], *keys: str) -> str | None:
        value = self._value(row, *keys)
        if value is None:
            return None
        text = str(value).strip()
        return mapping.get(self._normalize_header(text), text)

    def _required(self, row: dict[str, Any], *keys: str) -> Any:
        value = self._value(row, *keys)
        if value is None:
            raise ValueError(f"Campo requerido faltante: {keys[0]}")
        return value

    def _uuid_value(self, row: dict[str, Any], *keys: str) -> str:
        value = self._value(row, *keys)
        return str(value) if value is not None else str(uuid4())

    def _bool_value(self, row: dict[str, Any], *keys: str, default: bool = False) -> bool:
        value = self._value(row, *keys)
        if value is None:
            return default
        if isinstance(value, bool):
            return value
        if isinstance(value, (int, float)):
            return bool(value)
        normalized = str(value).strip().lower()
        return normalized in {"1", "si", "sí", "true", "verdadero", "yes", "y"}

    def _int_value(self, row: dict[str, Any], *keys: str, default: int | None = None) -> int | None:
        value = self._value(row, *keys)
        if value is None:
            return default
        try:
            return int(Decimal(self._numeric_text(value)))
        except InvalidOperation as exc:
            raise ValueError(f"Valor numerico invalido: {value}") from exc

    def _decimal_value(self, row: dict[str, Any], *keys: str) -> Decimal | None:
        value = self._value(row, *keys)
        if value is None:
            return None
        if isinstance(value, Decimal):
            return value
        try:
            return Decimal(self._numeric_text(value))
        except InvalidOperation as exc:
            raise ValueError(f"Valor numerico invalido: {value}") from exc

    def _seniority_months_value(self, row: dict[str, Any]) -> int | None:
        months = self._int_value(row, "antiguedad_meses", "seniority_months")
        if months is not None:
            return months
        years = self._int_value(row, "antiguedad_anos", "antiguedad_anios", "seniority_years")
        return years * 12 if years is not None else None

    def _numeric_text(self, value: Any) -> str:
        text = str(value).strip()
        text = re.sub(r"[^\d,.\-+eE]", "", text)
        if not text or not re.search(r"\d", text):
            raise ValueError(f"Valor numerico invalido: {value}")
        if "e" in text.lower():
            return text.replace(",", ".")
        if "," in text and "." in text:
            if text.rfind(",") > text.rfind("."):
                text = text.replace(".", "").replace(",", ".")
            else:
                text = text.replace(",", "")
        elif "," in text:
            text = text.replace(",", ".")
        return text

    def _date_value(self, row: dict[str, Any], *keys: str, required: bool = False) -> date | None:
        value = self._value(row, *keys)
        if value is None:
            if required:
                raise ValueError(f"Campo requerido faltante: {keys[0]}")
            return None
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        text = str(value).strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
        return date.fromisoformat(text)

    def _datetime_value(self, row: dict[str, Any], *keys: str) -> datetime | None:
        value = self._value(row, *keys)
        if value is None:
            return None
        if isinstance(value, datetime):
            return value
        if isinstance(value, date):
            return datetime.combine(value, datetime.min.time())
        text = str(value).strip()
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y"):
            try:
                return datetime.strptime(text, fmt)
            except ValueError:
                continue
        return datetime.fromisoformat(text)

    def _clean_code(self, value: Any) -> str | None:
        text = self._clean_text(value)
        if text is None:
            return None
        return text.upper()

    def _is_uuid(self, value: Any) -> bool:
        try:
            UUID(str(value))
        except (TypeError, ValueError):
            return False
        return True

    def _stable_id(self, dataset: str, code: str | None) -> str:
        if not code:
            return str(uuid4())
        return str(uuid5(NAMESPACE_URL, f"asur-antifraude:{dataset}:{code}"))
